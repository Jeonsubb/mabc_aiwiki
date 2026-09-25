# -*- coding: utf-8 -*-
"""Nodus PRD v0.2 → PPTX (16:9, 단색 + 초록 로고)"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

OUT = os.path.expanduser("~/Downloads/Nodus-PRD-v0.2.pptx")
ASSETS = os.path.expanduser("~/mabc_aiwiki-ui-integration/docs/poster-assets/")

INK = RGBColor(0x14, 0x18, 0x1C); BODY = RGBColor(0x14, 0x18, 0x1C); MUTE = RGBColor(0x14, 0x18, 0x1C)
FAINT = RGBColor(0x14, 0x18, 0x1C); LINE = RGBColor(0xCF, 0xD5, 0xDA); PANEL = RGBColor(0xF3, 0xF5, 0xF6)
PANEL2 = RGBColor(0xEC, 0xEF, 0xF1); WHITE = RGBColor(0xFF, 0xFF, 0xFF); GREEN = RGBColor(0x2F, 0x8F, 0x7A)
FONT = "Apple SD Gothic Neo"

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333

def txt(slide, x, y, w, h, text, size=14, bold=False, color=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, line=1.15):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return tb

def bullets(slide, x, y, w, h, items, size=13, color=BODY, gap=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.15
        if isinstance(it, tuple):
            r = p.add_run(); r.text = "•  " + it[0]; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = FONT
            r2 = p.add_run(); r2.text = " " + it[1]; r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = FONT
        else:
            r = p.add_run(); r.text = "•  " + it; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT
    return tb

def rect(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.06
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp

def logo(slide, x, y, s=0.42, color=GREEN):
    """세 점을 잇는 노드 로고 (초록)"""
    def dot(cx, cy, r):
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + cx*s - r*s), Inches(y + cy*s - r*s), Inches(2*r*s), Inches(2*r*s))
        d.fill.background(); d.line.color.rgb = color; d.line.width = Pt(2.2); d.shadow.inherit = False
    def seg(x1, y1, x2, y2):
        c = slide.shapes.add_connector(1, Inches(x + x1*s), Inches(y + y1*s), Inches(x + x2*s), Inches(y + y2*s))
        c.line.color.rgb = color; c.line.width = Pt(2.2)
    dot(0.33, 0.40, 0.095); dot(0.67, 0.29, 0.095); dot(0.58, 0.65, 0.095)
    seg(0.42, 0.41, 0.58, 0.32); seg(0.67, 0.38, 0.62, 0.57)

def header(slide, num, title, sub=None):
    logo(slide, 0.55, 0.42, s=0.5)
    txt(slide, 1.15, 0.38, 0.9, 0.5, num, size=12, bold=True, color=MUTE, font="Menlo")
    txt(slide, 1.15, 0.62, 10.5, 0.7, title, size=26, bold=True, color=INK)
    if sub: txt(slide, 1.17, 1.28, 11.5, 0.5, sub, size=13, color=MUTE)
    ln = slide.shapes.add_connector(1, Inches(0.6), Inches(1.78), Inches(W-0.6), Inches(1.78)); ln.line.color.rgb = INK; ln.line.width = Pt(1.25)
    # footer
    txt(slide, 0.6, 7.05, 6, 0.3, "Nodus · PRD v0.2 · 2026-09-16", size=9, color=FAINT)
    txt(slide, W-3.6, 7.05, 3.0, 0.3, "MABC 2026 결선 · Let's do well", size=9, color=FAINT, align=PP_ALIGN.RIGHT)

def table(slide, x, y, w, rows, col_w=None, size=11, head_fill=PANEL2, row_h=0.42):
    nrows, ncols = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h*nrows)).table
    if col_w:
        for i, cw in enumerate(col_w): gt.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = gt.cell(r, c); cell.text = ""
            cell.margin_left = cell.margin_right = Inches(0.08); cell.margin_top = cell.margin_bottom = Inches(0.05)
            p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; run = p.add_run(); run.text = str(val)
            run.font.size = Pt(size); run.font.name = FONT
            run.font.bold = (r == 0); run.font.color.rgb = INK if r == 0 else BODY
            cell.fill.solid(); cell.fill.fore_color.rgb = head_fill if r == 0 else (WHITE if r % 2 else RGBColor(0xFA, 0xFB, 0xFB))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return gt

def card(slide, x, y, w, h, eyebrow, title, body, fill=PANEL, title_size=15, body_size=11.5):
    rect(slide, x, y, w, h, fill=fill, line=LINE)
    if eyebrow: txt(slide, x+0.18, y+0.14, w-0.36, 0.3, eyebrow, size=9.5, bold=True, color=MUTE, font="Menlo")
    txt(slide, x+0.18, y+(0.42 if eyebrow else 0.16), w-0.36, 0.6, title, size=title_size, bold=True, color=INK)
    if body:
        if isinstance(body, list):
            bullets(slide, x+0.12, y+(0.98 if eyebrow else 0.72), w-0.3, h-1.0, body, size=body_size, gap=3)
        else:
            txt(slide, x+0.18, y+(0.98 if eyebrow else 0.72), w-0.36, h-1.0, body, size=body_size, color=BODY, line=1.25)

def flow(slide, y, steps, h=1.55, x0=0.6, gap=0.18):
    n = len(steps); w = (W - 2*x0 - gap*(n-1)) / n
    for i, (kind, t, d) in enumerate(steps):
        x = x0 + i*(w+gap)
        dark = kind == "human"
        rect(slide, x, y, w, h, fill=INK if dark else PANEL, line=None if dark else LINE)
        b = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+0.18), Inches(y+0.19), Inches(0.26), Inches(0.26))
        b.fill.solid(); b.fill.fore_color.rgb = WHITE if dark else INK; b.line.fill.background(); b.shadow.inherit = False
        tf = b.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(i+1); r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = INK if dark else WHITE; r.font.name = "Menlo"
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        txt(slide, x+0.5, y+0.12, w-0.62, 0.4, t, size=14, bold=True, color=WHITE if dark else INK)
        txt(slide, x+0.18, y+0.62, w-0.36, h-0.7, d, size=10.5, color=RGBColor(0xE6,0xE9,0xEB) if dark else BODY, line=1.2)
        if i < n-1:
            txt(slide, x+w-0.02, y+h/2-0.22, gap+0.06, 0.4, "›", size=18, color=FAINT, align=PP_ALIGN.CENTER)

# ───────────────────────── 1. 표지
s = prs.slides.add_slide(BLANK)
logo(s, 0.9, 0.85, s=1.0)
txt(s, 1.95, 0.9, 5, 0.9, "Nodus", size=34, bold=True, color=INK)
txt(s, 0.9, 2.2, 6.2, 3.2, ["AI 대화를", "승인 통제형", "개인 위키와", "생각 지도로", "바꾸는 서비스"], size=40, bold=True, color=INK, line=1.12)
hp = os.path.join(ASSETS, "01-home.png")
if os.path.exists(hp):
    pic = s.shapes.add_picture(hp, Inches(7.0), Inches(1.55), width=Inches(5.55)); pic.line.color.rgb = LINE; pic.line.width = Pt(0.75)
ln = s.shapes.add_connector(1, Inches(0.9), Inches(5.85), Inches(6.4), Inches(5.85)); ln.line.color.rgb = LINE
meta = [("팀", "Let's do well"), ("팀원", "이영민, 권윤재, 전동훈")]
for i, (k, v) in enumerate(meta):
    x = 0.9 + i*2.6
    txt(s, x, 6.0, 2.6, 0.3, k, size=11, bold=True, color=MUTE)
    txt(s, x, 6.3, 3.2, 0.4, v, size=14, color=INK)

# ───────────────────────── 2. 한눈에 보는 Nodus
s = prs.slides.add_slide(BLANK); header(s, "00", "한눈에 보는 Nodus", "보관 → 후보 생성 → 제안 → 사용자 확인 → 반영. 승인 게이트를 지나지 않은 내용은 위키에 쓰이지 않는다")
flow(s, 2.15, [
    ("auto", "대화 보관", "사용자가 에이전트에게 \"위키에 저장해\"라고 하면 MCP로 원문 구간이 전달 및 보관"),
    ("auto", "후보 생성", "Solar Pro 4 + ai-wiki 스킬이 기존 위키와 대조해 신규·갱신·연결 후보 생성"),
    ("auto", "제안 제시", "제안 이유, 원문 근거, 변경 전후 미리보기를 함께 제시"),
    ("human", "사용자 확인", "사용자가 근거를 보고 수락 또는 기각"),
    ("human", "위키 반영", "수락한 것만 위키·버전 이력·생각 지도에 반영"),
], h=1.9)
rect(s, 0.6, 4.4, 0.32, 0.22, fill=PANEL, line=LINE, radius=False); txt(s, 0.98, 4.33, 2, 0.35, "자동화 구간", size=10.5, color=MUTE)
rect(s, 2.3, 4.4, 0.32, 0.22, fill=INK, radius=False); txt(s, 2.68, 4.33, 2.5, 0.35, "사용자 승인 구간", size=10.5, color=MUTE)
card(s, 0.6, 4.95, 5.95, 1.85, "기능 1", "Solar Pro 4 기반 서비스 자체 에이전트", "대화·메모를 읽어 위키 노드 초안과 기존 위키에 대한 갱신·연결 제안을 생성합니다. ai-wiki 스킬(ai-wiki-SKILL.md)의 운영 규칙을 지침으로 삼습니다.")
card(s, 6.78, 4.95, 5.95, 1.85, "기능 2", "그래프 자동 연결·계층화", "대화를 노드와 관계로 잇고 주제→하위 주제→엔티티/아이디어로 조직합니다. 자동 연결은 후보로만 제시하고 실제 반영은 사용자 승인 아래 둡니다.")


# ───────────────────────── 2b. 아키텍처
from pptx.enum.dml import MSO_LINE
s = prs.slides.add_slide(BLANK); header(s, "00", "아키텍처", "사용자 요청 → 에이전트 → MCP → Nodus 서버 → Solar Pro 4 / PostgreSQL → Nodus 웹")
def abox(x, y, w, h, title, sub):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    txt(s, x, y+0.18, w, 0.45, title, size=15, bold=True, color=INK, align=PP_ALIGN.CENTER)
    txt(s, x, y+0.68, w, 0.6, sub, size=10.5, color=INK, align=PP_ALIGN.CENTER, line=1.2)
def dline(x1, y1, x2, y2):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor(0x8A,0x93,0x9B); c.line.width = Pt(1.25); c.line.dash_style = MSO_LINE.DASH
def pill(x, y, w, text):
    r = rect(s, x, y, w, 0.3, fill=RGBColor(0xF1,0xF3,0xF4), line=LINE); r.adjustments[0] = 0.5
    txt(s, x, y+0.02, w, 0.28, text, size=9.5, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# 상단 행
abox(0.9, 2.2, 2.4, 1.15, "사용자", "AI 대화 · 승인 결정")
def actor(cx, top, h=0.55):
    hd = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx-0.08), Inches(top), Inches(0.16), Inches(0.16)); hd.fill.background(); hd.line.color.rgb = INK; hd.line.width = Pt(1.5); hd.shadow.inherit = False
    for (x1,y1,x2,y2) in [(cx,top+0.16,cx,top+0.38),(cx-0.15,top+0.24,cx+0.15,top+0.24),(cx,top+0.38,cx-0.12,top+h),(cx,top+0.38,cx+0.12,top+h)]:
        c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); c.line.color.rgb = INK; c.line.width = Pt(1.5)
actor(1.25, 2.5)
abox(4.2, 2.2, 4.4, 1.15, "AI 에이전트 (MCP 클라이언트)", "Hermes Agent · Solar Pro 4 · \"위키에 저장해\" 요청")
# 하단 행
abox(0.9, 4.6, 3.3, 1.35, "Frontend", "React · Vite · three.js · Vercel")
abox(4.9, 4.6, 3.4, 1.35, "Backend", "Node.js · Express · TypeScript · Render")
abox(9.6, 4.6, 2.85, 1.35, "Solar Pro 4", "Upstage API · ai-wiki 스킬을 지침으로")
abox(4.9, 6.15, 3.4, 0.75, "PostgreSQL · Prisma", "")
# 연결선
dline(3.3, 2.775, 4.2, 2.775)
dline(2.1, 3.35, 2.1, 4.6); pill(1.65, 3.8, 0.9, "웹 UI")
dline(6.4, 3.35, 6.4, 3.95); dline(6.4, 3.95, 6.6, 3.95); dline(6.6, 3.95, 6.6, 4.6); pill(6.75, 3.8, 3.2, "MCP · submit_conversation · Bearer 토큰")
dline(4.2, 5.27, 4.9, 5.27); pill(4.1, 5.05, 0.9, "REST /api")
dline(8.3, 5.27, 9.6, 5.27); pill(8.45, 5.05, 1.0, "후보 · 관계")
dline(6.6, 5.95, 6.6, 6.15)

# ───────────────────────── 3. 개요 · 목표
s = prs.slides.add_slide(BLANK); header(s, "01", "개요 · 목표")
txt(s, 0.6, 2.0, 5.9, 0.4, "제품 설명", size=15, bold=True, color=INK)
txt(s, 0.6, 2.45, 5.9, 3.0, ["사용자가 AI와 주고받은 대화를 읽고, 개인화된 주제별 위키로 정리해 제공하는 웹 서비스입니다.", "", "에이전트가 생성한 결과는 위키에 바로 쓰지 않고 제안 형태로만 제시하며, 사용자가 수락한 내용만 개인 위키에 반영됩니다.", "", "개발자뿐 아니라 AI를 사용하는 모든 사용자를 대상으로 하며, 아이디어·결정·계획·배움·회고 같은 일상적 대화도 위키로 다룹니다."], size=12.5, color=BODY, line=1.3)
rect(s, 0.6, 5.55, 5.9, 1.25, fill=PANEL, line=LINE)
txt(s, 0.78, 5.65, 5.6, 1.1, ["자동과 승인의 구분", "후보 생성은 자동화 영역, 실제 반영은 사용자 승인 영역입니다. 자동 제안이 곧바로 위키에 반영되는 구조는 아니다."], size=11, color=BODY, line=1.25)
txt(s, 6.9, 2.0, 5.9, 0.4, "목표", size=15, bold=True, color=INK)
bullets(s, 6.9, 2.45, 5.85, 4.5, [
    "개발자뿐 아니라 일반 사용자도 부담 없이 쓸 수 있는 낮은 진입장벽과 친근한 경험",
    "외부 AI 대화를 MCP를 통해 사용자별로 안전하게 수신·보관",
    "Solar Pro 4가 ai-wiki 스킬로 신규 위키 생성, 기존 위키 갱신, 위키 간 연결을 제안",
    "사용자가 변경 내용과 원문 근거를 확인한 뒤 수락 또는 기각",
    "수락된 결과만 개인 위키와 생각 지도에 반영",
    "같은 주제의 후속 대화가 기존 위키의 갱신 제안으로 이어짐",
    "AutoSchema의 스키마 유도 방식을 참고해 노드의 개념적 역할과 관계 유형을 판단",
], size=12)

# ───────────────────────── 4. 범위
s = prs.slides.add_slide(BLANK); header(s, "01", "범위 (포함 / 제외)", "MVP에 포함하는 것과 의도적으로 제외한 것")
table(s, 0.6, 2.0, 12.13, [
    ["포함", "제외"],
    ["AI와 주고받은 대화를 주제별 위키로 정리", "단순 채팅 로그 보관만 하고 구조화하지 않는 방식"],
    ["새 대화 유입 시 기존 위키 대조 후 신규·갱신·연결 제안 (보수적 제안 + 근거)", "근거 없이 연관성만 나열하는 추천 스팸성 기능"],
    ["과거 아이디어와 최근 맥락 사이 관계를 근거와 함께 제시", "근거 표기 없는 자동 연관성 엔진"],
    ["MCP를 통한 사용자 요청 기반 대화 원문 수신·보관", "사용자 요청 없이 대화를 자동 수집하는 기능"],
    ["Solar Pro 4가 ai-wiki 스킬로 만드는 신규·갱신·연결 제안", "AI 생성 내용을 사용자 승인 없이 바로 위키에 반영"],
    ["변경 전후와 원문 근거를 확인한 뒤 수락·기각", "제안 내용을 사용자가 직접 편집하는 기능"],
    ["AutoSchema의 스키마 유도 방식을 참고한 관계 유형 추론", "AutoSchemaKG 전체 파이프라인·자동 스키마 DB 구축"],
    ["사용자가 승인한 위키 관계를 생각 지도에 표시", "승인되지 않은 관계를 자동 확정"],
    ["위키와 보관 대화에 대한 기본 키워드 검색", "개인화 랭킹·벡터 검색·그래프 경로 검색"],
], col_w=[6.065, 6.065], size=11, row_h=0.46)
for c in range(2):
    s.shapes[-1].table.cell(0, c).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# ───────────────────────── 5. 문제 정의
s = prs.slides.add_slide(BLANK); header(s, "02", "문제 정의")
txt(s, 0.6, 2.0, 12, 0.35, "배경", size=15, bold=True, color=INK)
bullets(s, 0.6, 2.4, 12, 1.2, ["AI를 자주 쓰는 사람들은 대화 내용이 쌓이지만, 나중에 다시 찾아보기 어렵습니다.", "기록이 흩어져 있으면 예전 아이디어가 지금 맥락과 연결되는 지점을 놓치기 쉽습니다.", "위키로 정리하면 꺼내 쓰기 좋지만, 사람이 매번 정리하기엔 부담이 큽니다."], size=12, gap=3)
cw = (12.13 - 0.4) / 3
card(s, 0.6, 3.75, cw, 1.55, "문제 1", "기록으로만 남는 대화", "AI와 주고받은 대화가 기록으로만 남고, 주제별 지식으로 다시 꺼내 쓰기 어렵습니다.")
card(s, 0.6 + cw + 0.2, 3.75, cw, 1.55, "문제 2", "연결되지 않는 새 대화", "새 대화가 들어와도 기존 기록과 연결되지 않아 과거 아이디어와의 관계가 표면화되지 않습니다.")
card(s, 0.6 + 2*(cw + 0.2), 3.75, cw, 1.55, "문제 3", "높은 진입장벽", "개발자 중심으로만 설계된 도구는 비개발자가 쓰기엔 진입장벽이 높습니다.")
txt(s, 0.6, 5.5, 12, 0.35, "기대 효과", size=15, bold=True, color=INK)
bullets(s, 0.6, 5.9, 6.0, 1.2, ["따로 정리하지 않아도 대화가 위키 형태로 남아 재탐색 비용이 낮아집니다.", "새 대화 유입 시 기존 위키와의 관계가 제안되어 위키가 진화하는 지식 기반이 됩니다."], size=11.5, gap=2)
bullets(s, 6.75, 5.9, 6.0, 1.2, ["변경 내용과 원문 근거를 확인한 뒤 수락 여부를 결정할 수 있습니다.", "승인된 관계가 생각 지도에 표시되어 지식 구조를 시각적으로 탐색할 수 있습니다."], size=11.5, gap=2)

# ───────────────────────── 6. 페르소나 · 여정
s = prs.slides.add_slide(BLANK); header(s, "03", "목표 사용자 · 사용자 여정")
cw = (12.13 - 0.4) / 3
card(s, 0.6, 2.0, cw, 2.35, "사용자 1", "AI로 아이디어를 정리하며 일하는 개인", ["블로그·공부·기획을 하는 비개발자. 기록이 흩어져 있어 예전 대화를 이어가려면 매번 찾고 다시 설명하는 데 힘이 드는 상황", ("니즈", "주제별 자동 정리, 예전 생각과 지금 생각의 연결"), ("페인", "기록할 시간은 부족하고 쌓아둔 대화는 다시 안 봄")], body_size=10.5)
card(s, 0.6 + cw + 0.2, 2.0, cw, 2.35, "사용자 2", "여러 주제를 AI와 대화하며 배우는 사용자", ["학습·취미·의사결정·회고 등 다양한 맥락에서 AI와 소통하는 사람. 주제가 늘어나면 어디서 무엇을 배웠는지 되짚기 어려워 같은 질문을 다시 하는 데 힘이 드는 상황", ("니즈", "아이디어·결정·배움 단위로 위키가 생기길 원함"), ("페인", "대화가 많아져도 구조가 없어 전체 그림이 안 보임")], body_size=10.5)
card(s, 0.6 + 2*(cw + 0.2), 2.0, cw, 2.35, "사용자 3 · 향후", "소규모 팀 / 커뮤니티", "여러 사람이 AI와 주고받은 맥락을 함께 보고 싶은 팀. 지금은 구체화하지 않고 TBD로 둡니다.", fill=WHITE, body_size=10.5)
table(s, 0.6, 4.6, 12.13, [
    ["터치포인트", "현재의 고통 지점", "Nodus가 개입하는 방식"],
    ["대화 발생", "기록이 흩어지고 끝나버림", "\"위키에 저장해\" 한마디로 MCP를 통해 원문을 보관"],
    ["위키 확인", "별도 정리 없으면 찾기 어려움", "Solar가 주제별 노드 초안을 만들고 사용자가 수락해 축적"],
    ["새 대화 유입", "기존 위키와 대조해 연결이 보이지 않음", "기존 위키와 대조해 갱신·연결 제안을 근거와 함께 제시"],
    ["과거 회상", "예전 아이디어가 다시 연결되어도 표면화되지 않음", "승인된 관계를 생각 지도에 표시하고 키워드로 위키·보관 대화를 검색"],
], col_w=[2.2, 4.2, 5.73], size=10.5, row_h=0.42)
_jt = s.shapes[-1].table
for _r in range(5):
    for _c in range(3 if _r == 0 else 2):
        _jt.cell(_r, _c).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# ───────────────────────── 7. 기능 개요
s = prs.slides.add_slide(BLANK); header(s, "04", "기능 요구사항 개요", "핵심 기능 F1–F5 (MVP 기준)")
table(s, 0.6, 2.0, 12.13, [
    ["ID", "기능", "핵심 출력", "우선순위"],
    ["F1", "대화 읽어서 위키 노드 만들기", "제목·요약·본문·주제·태그·분류를 가진 신규 위키 후보, 또는 기존 위키 갱신 제안", "P0"],
    ["F2", "제안 수락/기각", "수락 시 노드 생성·갱신·관계 저장, 기각 시 미반영, 결정 상태 유지", "P0"],
    ["F3", "새 대화 유입 시 기존 위키 대조 후 제안", "신규·갱신·연결 제안 + 이유 + 원문 근거 + 변경 전후 미리보기", "P0"],
    ["F4", "위키와 보관 대화 검색", "키워드가 포함된 위키 노드와 보관 대화, 상세·원본으로 이동", "P1"],
    ["F5", "Solar 기반 동적 관계 스키마 유도 및 연결 제안", "출발·도착 위키, 개념적 역할, 관계 유형, 판단 이유, 원문 근거", "P0"],
], col_w=[0.8, 4.2, 6.0, 1.13], size=11.5, row_h=0.62)
rect(s, 0.6, 6.0, 12.13, 0.85, fill=PANEL, line=LINE)
txt(s, 0.8, 6.1, 11.8, 0.7, "공통 원칙: 생성 결과는 위키에 바로 반영하지 않고 제안 상태로 저장합니다. 초반에는 가장 신뢰도 높은 제안만 보수적으로 보여주고, 사용자가 수락하기 전에 무엇이 바뀌는지 알 수 있어야 합니다.", size=11.5, color=BODY, line=1.3)

# ───────────────────────── 8. F1 · F2
s = prs.slides.add_slide(BLANK); header(s, "04", "F1 · F2: 위키 노드 생성과 제안 수락/기각", "우선순위 P0")
card(s, 0.6, 2.0, 5.95, 4.85, "F1 · P0", "대화 읽어서 위키 노드 만들기", [
    "MCP로 전달된 대화 원문을 Solar Pro 4가 ai-wiki 스킬로 분석해 주제별 위키 노드 후보를 만듭니다.",
    "기존 위키와 같은 주제라면 신규 노드보다 기존 위키 갱신을 우선합니다.",
    ("입력", "MCP로 전달된 대화 원문, 기존 위키 노드"),
    ("출력", "제목·요약·본문·주제·태그·분류를 포함한 신규 위키 후보 또는 갱신 제안"),
    ("제약", "생성 결과는 제안 상태로만 저장한다"),
], body_size=11.5)
card(s, 6.78, 2.0, 5.95, 4.85, "F2 · P0", "제안 수락 / 기각", [
    "사용자는 자동 생성된 제안의 내용과 근거를 검토한 뒤 수락하거나 기각합니다.",
    ("수락", "신규 제안은 위키 노드로 저장, 갱신 제안은 기존 위키와 버전 이력에 반영, 연결 제안은 위키 관계로 저장"),
    ("기각", "위키에 반영하지 않음"),
    ("유지", "결정 상태와 시각은 새로고침 후에도 유지"),
    ("제약", "제안 내용을 직접 편집하는 기능은 MVP 범위 밖"),
], body_size=11.5)

# ───────────────────────── 9. F3 · F4
s = prs.slides.add_slide(BLANK); header(s, "04", "F3 · F4: 대조 제안과 검색", "F3 우선순위 P0 · F4 우선순위 P1")
card(s, 0.6, 2.0, 5.95, 4.85, "F3 · P0", "새 대화 유입 시 기존 위키 대조 후 제안", [
    "새 대화가 들어오면 기존 위키를 검토해 신규 생성·갱신·연결이 필요한 내용을 제안합니다.",
    ("출력", "신규 위키 생성 제안, 기존 위키 갱신 제안, 위키 간 연결 제안, 제안 이유와 원문 근거, 갱신 제안의 변경 전후 미리보기"),
    ("제약", "가장 신뢰도 높은 제안만 보수적으로 제시. 수락 전에 무엇이 바뀌는지 보여줍니다. 결정과 결과는 새로고침·재접속 후에도 유지"),
], body_size=11.5)
card(s, 6.78, 2.0, 5.95, 4.85, "F4 · P1", "위키와 보관 대화 검색", [
    "사용자는 키워드로 승인된 위키와 서비스에 보관된 대화를 함께 검색할 수 있습니다.",
    ("입력", "사용자의 검색어"),
    ("출력", "검색어가 포함된 위키 노드, 검색어가 포함된 보관 대화, 결과에서 위키 상세 또는 원본 대화로 이동"),
    ("제약", "개인화 랭킹, 벡터 검색, 그래프 경로 검색은 MVP 범위에 포함하지 않는다"),
], body_size=11.5)

# ───────────────────────── 10. F5
s = prs.slides.add_slide(BLANK); header(s, "04", "F5: Solar 기반 동적 관계 스키마 유도 및 연결 제안", "우선순위 P0")
txt(s, 0.6, 2.0, 12.1, 0.9, "Solar Pro 4는 AutoSchema의 데이터 기반 스키마 유도 방식을 참고해, 기존 위키 노드와 새 대화에서 노드의 개념적 역할과 관계 유형을 추론합니다. 이를 바탕으로 근거가 있는 위키 연결 후보를 생성합니다.", size=12.5, color=BODY, line=1.3)
rect(s, 0.6, 2.95, 12.13, 0.6, fill=PANEL, line=LINE)
txt(s, 0.8, 3.05, 11.8, 0.45, "\"서로 관련된 위키가 있다면 어떤 관계인지와 왜 연결되는지 확인하고 직접 승인하고 싶다.\"", size=11.5, color=BODY)
card(s, 0.6, 3.8, 3.9, 3.0, None, "입력", ["새 대화 원문", "기존 위키 노드", "원문 세그먼트"], body_size=11.5)
card(s, 4.72, 3.8, 3.9, 3.0, None, "출력 / 결과", ["출발 위키와 도착 위키", "각 위키의 개념적 역할", "관계 유형", "관계 판단 이유", "원문 근거"], body_size=11.5)
card(s, 8.84, 3.8, 3.89, 3.0, None, "제약", ["공통 단어·태그가 같다는 이유만으로 연결하지 않음", "관계 의미와 원문 근거가 명확하지 않으면 제안하지 않음", "관계 후보는 승인 전 확정하지 않음", "수락한 관계만 생각 지도에 표시"], body_size=11)

# ───────────────────────── 11. 데이터 흐름
s = prs.slides.add_slide(BLANK); header(s, "04-1", "데이터 수신 · 저장 · 반영 흐름", "자동 제안이 곧바로 반영되는 구조가 아니라, 보관 → 후보 생성 → 제안 → 확인 → 반영 순서로 진행된다")
flow(s, 2.05, [
    ("auto", "전송 시작", "사용자의 명시적 저장 요청 → 에이전트가 MCP 도구 submit_conversation 호출"),
    ("auto", "보관", "원문 구간 + 맥락을 원본 보관 영역에 저장. 정리본과 분리"),
    ("auto", "후보 생성", "Solar Pro 4가 ai-wiki 스킬로 기존 위키와 대조해 신규·갱신·연결 제안 생성"),
    ("human", "사용자 확인", "변경 전후와 원문 근거를 보고 수락 / 기각"),
    ("human", "반영", "수락한 것만 위키·버전 이력·관계·생각 지도에 반영"),
], h=1.75)
txt(s, 0.6, 4.05, 6, 0.35, "전송 규칙", size=14, bold=True, color=INK)
bullets(s, 0.6, 4.42, 6.0, 2.5, [
    "사용자가 \"이 내용을 위키에 저장해줘\"라고 요청할 때만 전송합니다. 자동 스트리밍·주기 전송·대화 종료 자동 전송은 제외",
    "구간을 지정했으면 그 구간만, 지정하지 않았으면 현재 주제와 관련된 범위만 묶어 보낸다",
    "에이전트에게 실제로 보이는 대화만 보내고, 없는 대화를 추론해 만들지 않는다",
    "어떤 구간·주제인지 판단할 수 없을 때만 사용자에게 확인 질문을 한다",
], size=11, gap=3)
txt(s, 6.9, 4.05, 6, 0.35, "보관과 반영 원칙", size=14, bold=True, color=INK)
bullets(s, 6.9, 4.42, 5.85, 2.5, [
    "전달받은 기록은 먼저 보관합니다. 보관은 위키 반영과 분리된 원본성 기록이다",
    "보관 층은 둘: 서비스 수신 원본(MCP)과 사용자 추가 원본",
    "보관 기록은 Solar Pro 4가 ai-wiki 스킬로 후보를 만드는 입력으로만 쓴다",
    "제안은 적용 시 무엇이 바뀌는지(변경 전후)를 함께 보여주고, 결정과 결과는 재접속 후에도 남는다",
], size=11, gap=3)

# ───────────────────────── 12. 입력 계약
s = prs.slides.add_slide(BLANK); header(s, "04-1", "submit_conversation 입력 계약", "역할 구분 메시지 목록(messages) 또는 대화 텍스트(conversation_text)로 원문 구간을 전달한다")
table(s, 0.6, 2.0, 7.2, [
    ["필드", "필수", "설명"],
    ["session_id", "필수", "대화/세션을 식별하는 ID"],
    ["conversation_text", "선택", "전송할 대화 원문(구간). messages가 없을 때 사용"],
    ["messages[]", "선택", "role, content 필수 · record_id, timestamp 선택. 제공되면 원문의 1차 출처"],
    ["source", "선택", "전송 출처. mcp / agent / user_direct"],
    ["context", "선택", "주제·스킬·세션 유형 등 보충 맥락. 원문은 넣지 않음"],
], col_w=[1.9, 0.8, 4.5], size=10.5, row_h=0.5)
bullets(s, 0.6, 5.2, 7.2, 1.7, [
    "role과 content 중 하나라도 누락되면 도구 오류. 알 수 없는 id·시각은 추측해 넣지 않는다",
    "messages와 conversation_text가 함께 오면 정규화 후 비교해, 다르면 도구 오류로 처리합니다",
    "기존 conversation_text 기반 호출은 계속 지원한다",
], size=10.5, gap=3)
rect(s, 8.05, 2.0, 4.68, 4.85, fill=RGBColor(0xF4,0xF6,0xF7), line=LINE)
txt(s, 8.2, 2.1, 4.4, 4.7, [
 '{', '  "session_id": "sess-12345",', '  "source": "agent",', '  "messages": [', '    { "role": "user",', '      "content": "동네 산책 사진으로 사진집을', '                  만들까 고민 중이야.",', '      "record_id": "msg-001",', '      "timestamp": "2026-09-15T10:00:00+09:00" },', '    { "role": "assistant",', '      "content": "사진집은 인쇄와 전자책 두 가지', '                  형식이 있어요." }', '  ],', '  "context": { "topic": "사진집 기획",', '               "skill": "ai-wiki" }', '}'
], size=9.5, color=INK, font="Menlo", line=1.2)

# ───────────────────────── 13. 저장 구조 · 엔티티
s = prs.slides.add_slide(BLANK); header(s, "07", "데이터 모델 — 저장 구조와 핵심 엔티티", "원본 보관과 위키 정리본을 분리하고, 제안→승인→반영을 별도 상태·이력으로 추적한다")
rect(s, 0.6, 2.0, 5.5, 2.35, fill=WHITE, line=LINE)
txt(s, 0.8, 2.1, 5.1, 0.4, "원본 보관 영역", size=14, bold=True, color=INK)
bullets(s, 0.75, 2.5, 5.2, 1.8, ["MCP로 전달받은 대화 구간 + 맥락", "사용자가 직접 덧붙인 원본", "입력 당시의 기록을 그대로 보존", "세그먼트 단위로 근거 참조 가능"], size=11, gap=2)
txt(s, 6.15, 2.85, 0.9, 0.6, "→", size=28, color=FAINT, align=PP_ALIGN.CENTER)
rect(s, 7.1, 2.0, 5.63, 2.35, fill=RGBColor(0xF1,0xF3,0xF4), line=INK)
txt(s, 7.3, 2.1, 5.2, 0.4, "위키 정리본 영역", size=14, bold=True, color=INK)
bullets(s, 7.25, 2.5, 5.35, 1.8, ["주제/엔티티 단위 노드(제목·요약·본문·주제·태그·분류)", "수락된 갱신은 기존 내용을 보존하며 새 버전으로 누적", "수락된 연결은 관계로 저장되어 생각 지도에 표시", "어떤 원문이 어떤 노드에 반영됐는지 매핑"], size=11, gap=2)
table(s, 0.6, 4.6, 12.13, [
    ["엔티티", "역할"],
    ["User · McpCredential", "모든 데이터의 격리 단위 · 에이전트가 사용자 계정으로 제출할 때 쓰는 Bearer 토큰"],
    ["Record · ConversationSegment", "보관된 대화 원문과 맥락 · 원문을 근거 단위로 나눈 조각(제안 근거가 가리키는 대상)"],
    ["WikiNode · NodeVersion", "주제/엔티티 단위 정리본 · 갱신 이력"],
    ["Proposal · ProposalEvidence", "신규·갱신·연결 후보, 변경 전/후, 상태(제안됨·승인됨·기각됨·반영됨), 결정 시각 · 원문 인용"],
    ["NodeRelationship · RecordToNode", "수락된 연결에서 생긴 관계(유형·근거) · 어떤 원문 구간이 어떤 노드에 반영됐는지 추적"],
], col_w=[3.4, 8.73], size=10.5, row_h=0.38)

# ───────────────────────── 14. API
s = prs.slides.add_slide(BLANK); header(s, "07", "주요 API 엔드포인트", "구현 기준")
table(s, 0.6, 2.0, 12.13, [
    ["메서드", "경로", "설명"],
    ["MCP", "/mcp · submit_conversation", "에이전트가 사용자 토큰으로 대화 원문 제출 (Streamable HTTP)"],
    ["POST / GET", "/api/credentials", "MCP 연결 토큰 발급·목록·폐기"],
    ["GET / POST", "/api/records", "보관된 대화 목록 조회 · 원문 등록 후 후보 생성"],
    ["GET", "/api/proposals, /api/proposals/:id", "제안 목록·상세 (근거, 변경 전/후 포함)"],
    ["POST", "/api/proposals/:id/decide", "수락/기각. 수락 시 노드 생성·갱신 또는 관계 저장"],
    ["GET", "/api/nodes, /api/nodes/:id", "승인된 위키 노드 목록·상세"],
    ["GET", "/api/graph", "생각 지도용 노드와 승인된 관계(edges)"],
    ["GET", "/api/search/wiki-nodes, /api/search/chat-messages", "키워드 검색 (위키 노드 · 보관 대화)"],
], col_w=[1.5, 4.6, 6.03], size=11, row_h=0.5)

# ───────────────────────── 15. 주요 화면
s = prs.slides.add_slide(BLANK); header(s, "06", "주요 화면", "실제 서비스 화면 (2026-09-16 기준)")
shots = [("01-home.png", "생각 지도 (홈)", "승인된 노드와 관계를 그래프로 탐색. 노드 선택 시 요약·본문·대화로 이어가기"),
         ("03-proposal-detail.png", "제안 상세", "변경 전/후 비교와 원문 근거를 확인한 뒤 수락 또는 기각"),
         ("02-proposals.png", "제안 목록", "신규·갱신·연결 제안을 유형·상태 배지와 함께 나열"),
         ("04-search.png", "검색", "위키 노드와 보관 대화를 한 번에 검색, 연결된 위키를 근거로 표시")]
cw = (12.13 - 0.6) / 4
for i, (f, t, d) in enumerate(shots):
    x = 0.6 + i*(cw+0.2)
    p = os.path.join(ASSETS, f)
    if os.path.exists(p):
        pic = s.shapes.add_picture(p, Inches(x), Inches(2.0), width=Inches(cw))
        pic.line.color.rgb = LINE; pic.line.width = Pt(0.75)
    txt(s, x, 2.0 + cw*0.625 + 0.15, cw, 0.4, t, size=13, bold=True, color=INK)
    txt(s, x, 2.0 + cw*0.625 + 0.5, cw, 1.2, d, size=10.5, color=MUTE, line=1.25)

# ───────────────────────── 16. 비기능 · UX 원칙
s = prs.slides.add_slide(BLANK); header(s, "05 · 06", "비기능 요구사항 · 디자인 원칙", "")
table(s, 0.6, 2.0, 6.2, [
    ["항목", "요구사항"],
    ["성능", "위키 열람·제안 제시는 기다림을 느끼지 않는 수준. 파싱·요약·대조는 비동기 처리 검토"],
    ["보안", "저장·전송 시 보호, 사용자별 데이터 격리, MCP는 사용자별 Bearer 토큰 인증"],
    ["가용성", "핵심 위키 열람/제안은 안정적으로 제공"],
    ["확장성", "노드와 관계가 늘어나도 탐색/제안이 무너지지 않는 구조"],
    ["접근성", "비개발자도 읽기·탐색·제안을 이해하기 쉬워야 함"],
    ["로깅", "제안 수락/기각, 위키 조회 등 핵심 행동 로그"],
], col_w=[1.2, 5.0], size=10.5, row_h=0.55)
txt(s, 7.1, 2.0, 5.6, 0.4, "디자인 원칙", size=15, bold=True, color=INK)
bullets(s, 7.1, 2.45, 5.65, 4.4, [
    ("낮은 진입장벽.", "AI를 쓰는 일반 사용자도 처음 보고 이해할 수 있어야 한다"),
    ("명료성.", "신규/갱신/연결 같은 개념도 짧고 직관적으로 표현한다"),
    ("통제감.", "반영 여부는 항상 사용자가 수락/기각으로 결정한다"),
    ("자동과 승인의 분리.", "후보 생성은 자동화 구간, 위키 반영은 승인 구간이다"),
    ("근거 투명성.", "관계/제안을 제시할 때 왜 그런지 근거를 함께 보여준다"),
    ("과도한 자동화 회피.", "신뢰 높은 제안만 보수적으로 보여준다"),
], size=11.5, gap=5)

# ───────────────────────── 17. 성공 기준 · 리스크
s = prs.slides.add_slide(BLANK); header(s, "08 · 09", "MVP 성공 기준 · 리스크", "")
txt(s, 0.6, 2.0, 5.9, 0.4, "MVP 성공 기준", size=15, bold=True, color=INK)
tb = s.shapes.add_textbox(Inches(0.6), Inches(2.45), Inches(5.9), Inches(4.4)); tf = tb.text_frame; tf.word_wrap = True
for i, it in enumerate(["MCP로 제출한 대화가 사용자 계정의 유입 기록에 저장된다", "제출된 대화에서 최소 한 개 이상의 신규·갱신·연결 제안을 생성할 수 있다", "신규 제안을 수락하면 내 위키에 노드가 생성된다", "같은 주제의 후속 대화가 기존 위키 갱신 제안으로 이어진다", "연결 제안을 수락하면 생각 지도에 두 위키의 관계선이 표시된다", "기각한 제안은 위키에 반영되지 않는다", "새로고침 후에도 위키와 결정 상태가 유지된다"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(6)
    r = p.add_run(); r.text = "✓  "; r.font.bold = True; r.font.size = Pt(12.5); r.font.color.rgb = GREEN; r.font.name = FONT
    r2 = p.add_run(); r2.text = it; r2.font.size = Pt(12); r2.font.color.rgb = BODY; r2.font.name = FONT
table(s, 6.9, 2.0, 5.83, [
    ["리스크", "완화 방안"],
    ["제안 품질이 낮으면 스팸처럼 느껴질 수 있음", "초반 보수적 제안, 근거 표기, 사용자 피드백 반영"],
    ["자동 정리가 통제감을 약화시킬 수 있음", "모든 반영을 수락/기각 승인 뒤에만 수행"],
    ["대화가 적으면 위키/제안이 빈약할 수 있음", "빈 상태 UX, 점진적 적립 경험 설계"],
    ["비개발자에게 개념이 어려울 수 있음", "친근한 표현, 짧은 설명, 점진적 노출"],
], col_w=[3.0, 2.83], size=10.5, row_h=0.7)
txt(s, 6.9, 5.7, 5.83, 1.0, "외부 의존성: Solar Pro 4 API(Upstage), MCP 클라이언트를 가진 에이전트(Hermes 등), PostgreSQL, 프론트/백엔드 호스팅", size=10.5, color=MUTE, line=1.3)

# ───────────────────────── 18. 일정 · 변경 이력
s = prs.slides.add_slide(BLANK); header(s, "10", "일정", "")
table(s, 0.6, 2.0, 12.13, [
    ["단계", "기간", "산출물 / 완료 기준"],
    ["리서치 / 설계", "예선 ~ 2026-09-11", "핵심 사용자 유형, 위키 단위, 제안/관계 제시 방식, 데이터 모델 초안, PRD 0.1"],
    ["MVP 개발", "2026-09-12 ~ 09-16", "MCP 수신·보관, 대화→노드 생성, 대조 제안(신규·갱신·연결), 관계 제시+근거, 조회/탐색/검색, 수락/기각"],
    ["내부 테스트", "2026-09-16", "실제 대화 데이터로 위키·제안 품질 확인, 흐름 점검"],
    ["결선 제출 / 시연", "2026-09-16 ~", "핵심 흐름이 안정적으로 동작하고 비개발자도 쓸 수 있는 경험 확보. 제출 후 배포본 유지"],
], col_w=[2.2, 2.4, 7.53], size=11, row_h=0.5)


prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
